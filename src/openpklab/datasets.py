"""Strict parsers. Parsing never changes stored originals or evaluates formulas."""
from __future__ import annotations
import io
import csv
import warnings
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image

MAX_CELLS = 100_000


def read_table(filename: str, data: bytes) -> pd.DataFrame:
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        header = next(csv.reader(io.StringIO(data.decode("utf-8-sig"))), [])
        if not header or any(not c.strip() for c in header) or len(set(header)) != len(header):
            raise ValueError("Column headers must be nonempty and unique")
        frame = pd.read_csv(io.BytesIO(data), encoding="utf-8-sig", nrows=10_001)
    elif ext == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds 30 MiB")
            if any("vbaProject" in i.filename or "externalLinks/" in i.filename for i in archive.infolist()):
                raise ValueError("Macros and external workbook links are unsupported")
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) != 1:
                raise ValueError("Provide one worksheet per upload to avoid ambiguous table selection")
            sheet = book.active
            rows = []
            for row in sheet.iter_rows():
                if len(rows) > 10_000 or len(row) > 1000:
                    raise ValueError("Workbook exceeds row or column limit")
                if any(c.data_type == "f" for c in row):
                    raise ValueError("Formula workbooks are unsupported; export measured values")
                rows.append([c.value for c in row])
            if not rows:
                raise ValueError("Empty worksheet")
            if any(c is None for c in rows[0]) or len(set(rows[0])) != len(rows[0]):
                raise ValueError("Column headers must be nonempty and unique")
            frame = pd.DataFrame(rows[1:], columns=[str(c) for c in rows[0]])
        finally:
            book.close()
    else:
        raise ValueError("Expected CSV or XLSX")
    if frame.empty or frame.size > MAX_CELLS or len(frame) > 10_000 or len(frame.columns) > 1000:
        raise ValueError("Table must have 1–10000 rows, at most 1000 columns and 100000 cells")
    return frame


def validate_file(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    try:
        if ext in {".csv", ".xlsx"}:
            read_table(filename, data)
            return "table"
        if ext in {".png", ".jpg", ".jpeg"}:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(data)) as im:
                    if im.format not in {"PNG", "JPEG"} or im.width * im.height > 20_000_000:
                        raise ValueError("Expected PNG/JPEG, at most 20 megapixels")
                    if (ext == ".png") != (im.format == "PNG"):
                        raise ValueError("Image extension does not match its actual format")
                    im.verify()
            return "figure"
    except (ValueError, OSError, csv.Error, zipfile.BadZipFile, pd.errors.ParserError, Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
        raise ValueError(f"Invalid scientific file: {exc}") from exc
    raise ValueError("Supported file types: CSV, XLSX, PNG, JPEG. Documents and executable formats are unsupported.")


def quality(frame: pd.DataFrame) -> dict:
    findings = []
    missing = {str(c): int(frame[c].isna().sum()) for c in frame if frame[c].isna().any()}
    if missing:
        findings.append({"code": "missing_values", "priority": "high", "details": missing})
    duplicates = int(frame.duplicated().sum())
    if duplicates:
        findings.append({"code": "duplicate_rows", "priority": "medium", "count": duplicates})
    for column in frame:
        values = frame[column]
        if str(column).lower().endswith(("unit", "units")) and values.dropna().astype(str).nunique() > 1:
            findings.append({"code": "inconsistent_units", "priority": "high", "column": str(column), "units": sorted(values.dropna().astype(str).unique().tolist())})
        numeric = pd.to_numeric(values, errors="coerce")
        if numeric.notna().sum() >= max(1, values.notna().sum() / 2):
            invalid = int((values.notna() & numeric.isna()).sum())
            nonfinite = int((numeric.notna() & ~np.isfinite(numeric)).sum())
            if invalid or nonfinite:
                findings.append({"code": "invalid_numeric_entries", "priority": "high", "column": str(column), "count": invalid + nonfinite})
            if str(column).lower() in {"time", "concentration"} and (numeric < 0).any():
                findings.append({"code": "negative_pk_values", "priority": "high", "column": str(column), "count": int((numeric < 0).sum())})
            clean = numeric[np.isfinite(numeric)]
            if len(clean) >= 4:
                q1, q3 = clean.quantile([0.25, 0.75])
                iqr = q3 - q1
                out = clean[(clean < q1 - 1.5 * iqr) | (clean > q3 + 1.5 * iqr)]
                if len(out):
                    findings.append({"code": "outlier_screen", "priority": "low", "column": str(column), "row_numbers": [int(i) + 2 for i in out.index], "method": "1.5 IQR; screen only, no values removed"})
    return {"rows": len(frame), "columns": [str(c) for c in frame], "missing_values": missing, "duplicate_rows": duplicates, "completeness": float(frame.notna().to_numpy().mean()), "findings": findings, "original_modified": False}
